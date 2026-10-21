"""Portable regressions for the full workpack; set FISH_TEST_BUNDLE to its extracted directory."""
from pathlib import Path
import os, copy, hashlib, importlib.util, json, shutil, sqlite3, subprocess, sys, tempfile, unittest
from unittest.mock import patch

HERE=Path(__file__).resolve().parents[1]
PACKAGE=Path(os.environ.get('FISH_TEST_BUNDLE',str(HERE))).resolve()
sys.path[:0]=[str(PACKAGE/'toolkit'),str(PACKAGE/'toolkit/tools'),str(PACKAGE/'toolkit/skill/scripts')]
import install_bundle
from bundle_runtime import inspect
from analysis_checks import check,materialize
from report_model import sha,load,digest
from render_reviewed import render
from research import doctor,connect,init_project,register
from check_completion import check as completion
from openpyxl import load_workbook

def write(p,value):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(value,ensure_ascii=False,indent=2),'utf-8')
    return p

@unittest.skipUnless((PACKAGE/'toolkit/workspace.json').exists() and load(PACKAGE/'toolkit/workspace.json').get('database'), 'Legacy full-workpack checks require its bundled database; current software-only checks are in test_portable_install.py')
class RevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name);self.project=self.root/'课题';self.home=self.root/'workbuddy'
    def tearDown(self): self.temp.cleanup()
    def install(self,**kwargs): return install_bundle.install(self.project,workbuddy_home=self.home,**kwargs)
    def command(self,*args):
        return subprocess.run([sys.executable,'-X','utf8','-B',str(PACKAGE/'toolkit/tools/research.py'),*args],capture_output=True,text=True,encoding='utf8')

    def test_install_all_components_and_idempotence(self):
        r=self.install(client='workbuddy');self.assertEqual(r['errors'],[])
        self.assertEqual(r['modified_files'],[])
        self.assertEqual(r['bundle_version'],'1.5.0')
        self.assertEqual(self.install(client='workbuddy')['status'],'already_installed')
        d=doctor(self.project);self.assertTrue(d['ready_for_local_queries']);self.assertEqual(d['errors'],[])
        c,_,_=connect(self.project)
        try: self.assertEqual(d['counts']['proteins'],c.execute('SELECT count(*) FROM proteins').fetchone()[0])
        finally: c.close()

    def test_both_clients_are_idempotent(self):
        self.install(client='both')
        self.assertEqual(self.install(client='both')['status'],'already_installed')
        for name in ('.agents','.claude'):
            self.assertTrue((self.project/name/'skills/fish-toxicology-research/SKILL.md').exists())

    def test_discovery_mismatch_and_tool_tamper_detected(self):
        self.install(client='workbuddy')
        discovery=self.home/'skills/fish-toxicology-research/SKILL.md'
        discovery.write_text(discovery.read_text('utf8').replace('1.5.0','1.3.0'),'utf8')
        self.assertTrue(inspect(self.project/'.fish-toxicology')['errors'])
        tool=self.project/'.fish-toxicology/tools/research.py';tool.write_text('# modified','utf8')
        self.assertIn('tools/research.py',inspect(self.project/'.fish-toxicology')['modified_files'])

    def test_update_preserves_research_custom_profile_and_database(self):
        self.install(client='workbuddy')
        original=self.project/'research/personal.json';write(original,{'unchanged':True})
        profile=self.project/'profiles/custom.json';write(profile,{'theme':'custom'})
        config=self.project/'.fish-toxicology/workspace.json';c=load(config);c['custom_setting']='retained';write(config,c)
        tool=self.project/'.fish-toxicology/tools/research.py';tool.write_text('# local customization','utf8')
        with self.assertRaisesRegex(ValueError,'修改'): self.install(client='workbuddy',update=True)
        r=self.install(client='workbuddy',update=True,replace_managed=True)
        self.assertEqual(load(config)['database'],c['database']);self.assertEqual(load(config)['custom_setting'],'retained')
        self.assertEqual(load(original),{'unchanged':True});self.assertEqual(load(profile),{'theme':'custom'})
        index=load(Path(r['backup'])/'index.json')
        old=next(x for x in index if x['original']==str(tool))
        self.assertEqual((Path(r['backup'])/old['backup']).read_text('utf8'),'# local customization')

    def test_legacy_upgrade_on_copy(self):
        self.install(client='portable')
        record=self.project/'.fish-toxicology/install-record.json'
        current=load(record)
        legacy={'version':'1.2.1','clients':['portable'],
                'files':{Path(p).relative_to(self.project).as_posix():h for p,h in current['managed_files'].items()}}
        write(record,legacy)
        before=load(self.project/'.fish-toxicology/workspace.json')
        r=self.install(client='workbuddy',update=True,replace_managed=True)
        self.assertEqual(r['errors'],[]);self.assertEqual(r['bundle_version'],'1.5.0')
        self.assertEqual(load(self.project/'.fish-toxicology/workspace.json')['database'],before['database'])

    def test_failed_post_install_check_restores_files(self):
        self.install(client='workbuddy')
        before={p:p.read_bytes() for p in (self.project/'.fish-toxicology/tools').glob('*.py')}
        real=install_bundle.inspect
        def fail_at_project(p,*a,**kw):
            result=real(p,*a,**kw)
            if Path(p).name=='.fish-toxicology': result['errors']=['injected failure']
            return result
        with patch.object(install_bundle,'inspect',side_effect=fail_at_project):
            with self.assertRaisesRegex(ValueError,'injected'): self.install(update=True)
        self.assertTrue(all(p.read_bytes()==b for p,b in before.items()))

    def test_taxonomy_number_and_empty_query_fallback(self):
        right=self.command('taxonomy','Oryzias melastigma','--taxid','30732')
        wrong=self.command('taxonomy','Oryzias melastigma','--taxid','299683')
        self.assertEqual(right.returncode,0);self.assertNotEqual(wrong.returncode,0)
        empty=json.loads(self.command('proteins','--species','Oryzias melastigma','--term','nonexistent_test_term').stdout)
        all_species=json.loads(self.command('proteins','--species','Oryzias melastigma').stdout)
        self.assertEqual(empty['status'],'no_local_match');self.assertEqual(all_species['count'],102)

    def test_database_read_only_and_expression_unchanged(self):
        self.install();c,_,_=connect(self.project)
        try:
            with self.assertRaises(sqlite3.OperationalError): c.execute('CREATE TABLE injected_test(id)')
        finally: c.close()
        self.assertEqual(sha(PACKAGE/'toolkit/skill/references/expression.md'),'86baf14e28d56edd5a934bfb87b6b78f3253de80e7f9c6a79370a0a5371d1451')

    def test_cross_system_dose_claim_and_scope(self):
        d={'chemicals':[{'id':'A'}],'expected_chemical_ids':['A','B'],
           'experiments':[{'id':'X','object':'cell one','endpoint':'protein','duration':'24 h','protocol':'P'},
                          {'id':'Y','object':'cell two','endpoint':'transcript','duration':'48 h','protocol':'P'}],
           'comparisons':[{'id':'dose','type':'same_system_dose_response','experiment_ids':['X','Y']}]}
        errors=check(d);self.assertEqual(len(errors),4)
        d['comparisons'][0]['type']='cross_system_comparison';d['expected_chemical_ids']=['A']
        self.assertEqual(check(d),[])

    def test_recompute_geomean_and_references(self):
        d={'chemicals':[{'id':'A','prediction':'revised'}], 'paragraphs':{'x':[{'$ref':'chemicals/A/prediction'}]},
           'computations':[{'id':'ER','value':10}],
           'numeric_checks':[{'id':'geomean','operation':'geometric_mean','inputs':[4,15],'result_ref':'computations/ER/value'}]}
        self.assertTrue(check(d));d['computations'][0]['value']=(4*15)**.5
        self.assertEqual(check(d),[]);self.assertEqual(materialize(d)['paragraphs']['x'],['revised'])
        d['chemicals'][0]['prediction']={'$ref':'chemicals/A/prediction'}
        with self.assertRaisesRegex(ValueError,'Cyclic'): materialize(d)

    def test_detail_continuation_keeps_text(self):
        d=load(PACKAGE/'toolkit/skill/assets/report-example.json')
        text='这是用于验证实际长段落换行的内容。'*400
        d['chemicals'][0]['details']=[{'label':'长段落','text':text}]
        report=write(self.root/'report.json',d)
        spec=write(self.root/'presentation.json',{'analysis':'report.json','analysis_sha256':sha(report)})
        render(spec,self.root/'out',stage='preview')
        w=load_workbook(self.root/'out/analysis.xlsx');ws=w.worksheets[3]
        blocks=[(ws.cell(r,1).value,ws.cell(r,3).value) for r in range(4,ws.max_row+1)]
        pieces=[]
        for label,value in blocks:
            if label in ('长段落','长段落（续）'): pieces.append(value or '')
            elif pieces: break
        self.assertGreater(len(pieces),1);self.assertEqual(''.join(pieces),text)
        self.assertTrue(all((d.height or 0)<=409 for d in ws.row_dimensions.values()));w.close()

    def test_completion_rejects_project_config_as_review(self):
        d={'project_id':'P','analysis_version':'1','chemicals':[{'id':'A','prediction':'conditional'}]}
        a=write(self.root/'analysis.json',d)
        p={'status':'configured','research_question':'Q','comparison_conditions':'same','scope':{'candidates':['A']}}
        checks=['explicit_contradictions','negative_delivery_cases','expression_and_falsification','final_artifact_reading']
        common={'project_id':'P','analysis_version':'1','publication_sha256':digest(d),'reviewer':'local test fixture','findings':['fixture finding']}
        rev=write(self.root/'review.json',dict(common,record_type='publication_review',decision='accepted',chemical_ids=['A']))
        val=write(self.root/'validation.json',dict(common,record_type='validation_record',checks=checks))
        rec={'analysis':{'path':a.name,'sha256':sha(a)},
             'scope_dispositions':[{'category':'candidates','id':'A','status':'completed','contribution':'comparison','evidence_files':[a.name]}],
             'scientific_outputs':{k:a.name for k in ['candidate_predictions','risk_order_or_groups','representative_decisions','unresolved_items_and_impacts']},
             'content_review':{'status':'reviewed_with_stated_limits','reviewer':'local test fixture','record_file':rev.name,'sha256':sha(rev),'coverage':{'candidates':['A']}},
             'checks':{k:{'status':'passed','record_file':val.name,'sha256':sha(val)} for k in checks}}
        self.assertEqual(completion(p,rec,self.root)['errors'],[])
        bad=write(self.root/'project.json',p)
        rec['content_review'].update(record_file=bad.name,sha256=sha(bad))
        for c in rec['checks'].values(): c.update(record_file=bad.name,sha256=sha(bad))
        errors=completion(p,rec,self.root)['errors']
        self.assertTrue(any(e['code']=='UNBOUND_REVIEW' for e in errors))

    def test_blank_start_render_and_register_named_records(self):
        self.install()
        created=init_project(self.project,'round-one','same question')
        project=Path(created['project']);config=load(project/'project.json')
        self.assertEqual(config['scope']['candidates'],[])
        self.assertEqual(list((project/'analysis').iterdir()),[])
        d=load(PACKAGE/'toolkit/skill/assets/report-example.json')
        d.update(project_id='round-one',analysis_version='test-1')
        a=write(project/'analysis/report.json',d)
        rev=write(project/'review/content.json',{'record_type':'publication_review','project_id':'round-one',
            'analysis_version':'test-1','publication_sha256':digest(d),'chemical_ids':[c['id'] for c in d['chemicals']],
            'reviewer':'local test fixture','decision':'accepted','findings':['synthetic integration test']})
        spec=write(project/'analysis/presentation.json',{'analysis':a.name,'analysis_sha256':sha(a),
             'review':{'path':'../review/content.json','sha256':sha(rev)}})
        out=project/'deliverables/test-1';render(spec,out)
        m={'project_id':'round-one','analysis_version':'test-1','title':'fixture','analysis':'analysis/report.json','analysis_sha256':sha(a),
           'content_review':{'reviewer':'local test fixture','decision':'accepted','open_blockers':[],
               'candidate_ids':[c['id'] for c in d['chemicals']],'records':[{'path':'review/content.json','sha256':sha(rev)}]},
           'artifacts':[{'kind':'xlsx','path':'deliverables/test-1/analysis.xlsx','sha256':sha(out/'analysis.xlsx')},
                        {'kind':'interactive_html','path':'deliverables/test-1/index.html','sha256':sha(out/'index.html'),
                         'web_files':{p.relative_to(out).as_posix():sha(p) for p in out.rglob('*') if p.is_file()}}],
           'database_links':[{'candidate_id':c['id'],'chemical_id':None,'unmatched_reason':'synthetic fixture'} for c in d['chemicals']]}
        bad=copy.deepcopy(m);bad['content_review']['records']=[{'path':'project.json','sha256':sha(project/'project.json')}]
        f=write(project/'release.json',bad)
        with self.assertRaisesRegex(ValueError,'publication_review'): register(self.project,f)
        write(f,m);r=register(self.project,f);self.assertEqual(r['registered'],'round-one')
        self.assertEqual(len(load(self.project/'research/catalog.json')['analyses']),1)

if __name__=='__main__': unittest.main()
