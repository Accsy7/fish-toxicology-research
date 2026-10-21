"""Regression tests for real renderer failures and profile extensibility."""
from pathlib import Path
import copy
import json
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"toolkit/tools"))
from report_model import load,digest,sha,validate
from render_reviewed import render,check_review
from openpyxl import load_workbook


class RendererTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.base=Path(self.temp.name)
        self.profile=load(ROOT/"toolkit/skill/assets/profiles/accepted.json")
        self.profile["workbook"]["template"]=str(ROOT/"toolkit/skill/assets/workbook-template.xlsx")
        self.data=dict(schema_version="2.0",project_id="example",analysis_version="1",
            title="Example",context="Example conditions",chemicals=[
                dict(id="a",name="Compound A",group="G1",prediction="Conditional result",status="complete",basis="Observation",doubt="Open question",detail_link="Details",
                     details=[dict(label="Conditions",text="Line 1\nLine 2",source_ids=["S1"])])],
            sources=[dict(id="S1",title="Source",url="https://example.org",note="Example")],
            scope=[],representatives=[],experiments=[],computations=[],target_proteins=[],reference_proteins=[],unresolved=[],
            paragraphs={},panels=[dict(id="plot",type="matrix",section="mechanism",title="Test",columns=["Value"],rows=[dict(label="A",chemical_ids=["a"],values=["value"])])])

    def tearDown(self):
        self.temp.cleanup()

    def write(self,name,value):
        p=self.base/name;p.write_text(json.dumps(value,ensure_ascii=False),"utf-8");return p

    def make(self,data=None,profile=None):
        a=self.write("report.json",data or self.data)
        self.write("profile.json",profile or self.profile)
        return self.write("presentation.json",dict(analysis="report.json",analysis_sha256=sha(a),profile="profile.json"))

    def test_template_blocks_hyperlinks_and_freezes(self):
        p=self.make();out=self.base/"out";render(p,out,stage="preview")
        book=load_workbook(out/"analysis.xlsx")
        detail=book.worksheets[3]
        self.assertIn("A4:B4",[str(m) for m in detail.merged_cells.ranges])
        self.assertIn("C4:H4",[str(m) for m in detail.merged_cells.ranges])
        self.assertEqual(detail["C4"].value,"Conditional result")
        self.assertEqual(detail["C4"].font.name,"Microsoft YaHei")
        self.assertEqual(book.worksheets[7].freeze_panes,"C4")
        self.assertTrue(book.worksheets[2]["G5"].hyperlink.location.startswith("'03"))
        urls=[c for row in book.worksheets[4] for c in row if c.hyperlink and c.hyperlink.target]
        self.assertTrue(any(c.hyperlink.target=="https://example.org" for c in urls))
        self.assertTrue(all(not d.hidden for ws in book for d in ws.row_dimensions.values()))
        book.close()

    def test_custom_profile_can_add_fields_and_change_counts(self):
        p=copy.deepcopy(self.profile)
        p["workbook"]["sheets"]=[dict(name="Custom",title="Custom",template_sheet=0,
          columns_count=2,kind="table",collection="extra",header_row=3,header_sample=3,body_sample=4,
          columns=[dict(field="id",label="ID"),dict(field="response",label="Extra response")])]
        p["web"]["sections"]=[dict(id="extra",title="Extra",blocks=[dict(kind="table",id="x",title="Extra",collection="extra",columns=p["workbook"]["sheets"][0]["columns"])])]
        self.data["extra"]=[dict(id="custom",response=1.25)]
        out=self.base/"custom";render(self.make(profile=p),out,stage="preview")
        w=load_workbook(out/"analysis.xlsx");self.assertEqual(w.sheetnames,["Custom"]);self.assertEqual(w.active["B4"].value,1.25);w.close()

    def test_unknown_source_and_chemical_are_rejected(self):
        self.data["chemicals"][0]["details"][0]["source_ids"]=["S10"]
        self.data["panels"][0]["rows"][0]["chemical_ids"]=["missing"]
        errors=validate(self.data,self.profile)
        self.assertTrue(any("S10" in x for x in errors))
        self.assertTrue(any("missing" in x for x in errors))

    def test_blank_complete_record_is_rejected(self):
        self.data["chemicals"][0]["prediction"]=""
        self.assertTrue(any("no prediction" in x for x in validate(self.data,self.profile)))
        self.data["chemicals"][0]["status"]="unassessed"
        self.assertFalse(any("no prediction" in x for x in validate(self.data,self.profile)))

    def test_invalid_matrix_and_nonfinite_values(self):
        self.data["panels"][0]["rows"][0]["values"]=[]
        self.data["panels"].append(dict(id="number",type="interval",rows=[dict(value=float("nan"))]))
        errors=validate(self.data,self.profile)
        self.assertTrue(any("column" in x for x in errors))
        self.assertTrue(any("numeric" in x for x in errors))

    def test_review_bound_to_actual_project_and_content(self):
        r=dict(record_type="publication_review",project_id="example",analysis_version="1",
               publication_sha256=digest(self.data),decision="accepted",reviewer="Reviewer",
               chemical_ids=["a"],findings=["Read the analysis and checked its mapping."])
        f=self.write("review.json",r)
        check_review(dict(path=f.name,sha256=sha(f)),self.data,self.base)
        r["project_id"]="another-project";f=self.write("review.json",r)
        with self.assertRaisesRegex(ValueError,"project_id"):
            check_review(dict(path=f.name,sha256=sha(f)),self.data,self.base)
        f=self.write("review.json",{"project_id":"example"})
        with self.assertRaisesRegex(ValueError,"record_type"):
            check_review(dict(path=f.name,sha256=sha(f)),self.data,self.base)

    def test_analysis_hash_and_output_overwrite(self):
        spec=self.make();self.write("report.json",dict(changed=True))
        with self.assertRaisesRegex(ValueError,"hash"):
            render(spec,self.base/"bad",stage="preview")
        spec=self.make();render(spec,self.base/"out",stage="preview")
        with self.assertRaisesRegex(ValueError,"already exists"):
            render(spec,self.base/"out",stage="preview")

    def test_review_required_for_reviewed_output(self):
        with self.assertRaisesRegex(ValueError,"review"):
            render(self.make(),self.base/"out")

    def test_plain_source_strings_are_not_formulas(self):
        self.data["chemicals"][0]["prediction"]="=HYPERLINK(\"https://example.org\")"
        render(self.make(),self.base/"out",stage="preview")
        w=load_workbook(self.base/"out/analysis.xlsx")
        self.assertEqual(w.worksheets[3]["C4"].data_type,"s")
        w.close()

    def test_missing_link_label_is_filled_without_corrupting_readback(self):
        self.data['chemicals'][0].pop('detail_link')
        render(self.make(),self.base/'out',stage='preview')
        w=load_workbook(self.base/'out/analysis.xlsx')
        self.assertEqual(w.worksheets[2]['G5'].value,'查看')
        self.assertTrue(w.worksheets[2]['G5'].hyperlink.location.startswith("'03"));w.close()

    def test_scatter_log_domain_values_and_heatmap_width(self):
        self.data['panels']=[dict(id='s',type='scatter',y_scale='log',rows=[dict(x=1,value=0,size=-1)]),
                             dict(id='h',type='heatmap',columns=['A'],rows=[dict(values=[1,2])])]
        errors=validate(self.data,self.profile)
        self.assertTrue(any('numeric' in e for e in errors));self.assertTrue(any('column' in e for e in errors))


if __name__=="__main__":
    unittest.main()
