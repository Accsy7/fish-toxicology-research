"""Render an explicitly reviewed presentation using the accepted layout assets.

Scientific field mapping is supplied and reviewed by the analysis AI. This module
copies approved values; it never decides predictions, ranks, or evidence strength.
"""
from pathlib import Path
import copy
import hashlib
import html
import json
import math
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'vendor/python'))
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from research import load, sha, write


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                    separators=(',', ':')).encode('utf-8')).hexdigest()


def same_cell(actual, expected):
    # XLSX serialization rounds binary float tails to decimal precision.
    # No absolute tolerance: a missing/zero tiny effect must not pass as a value.
    if type(actual) in (int, float) and type(expected) in (int, float):
        return math.isclose(actual, expected, rel_tol=1e-14, abs_tol=0)
    return actual == expected or (actual is None and expected == '')


def validate(spec, base):
    errors = []
    a = load(base / spec['analysis'])
    if sha(base / spec['analysis']) != spec['analysis_sha256']:
        errors.append('正式分析与版面绑定不一致。')
    view = spec['publication']
    if a.get('publication') != view:
        errors.append('展示内容必须先保存到正式分析的publication，再复核和导出。')
    review = spec['publication_review']
    if review.get('decision') not in ('accepted', 'accepted_with_stated_limits') or not review.get('reviewer'):
        errors.append('缺少版面字段与正式分析的实际内容核对。')
    if review.get('publication_sha256') != digest(view):
        errors.append('待导出内容与已核对内容不一致。')
    record = review.get('record', {})
    if not record.get('path') or not record.get('sha256') or sha(base / record['path']) != record['sha256']:
        errors.append('内容核对记录缺失或已变化。')
    if len(view['sheets']) != 8 or [s['name'][:2] for s in view['sheets']] != [f'{i:02}' for i in range(8)]:
        errors.append('工作簿必须沿用00—07。')
    if len(view['sections']) != 4:
        errors.append('网页须保持四个科学问题入口。')
    ids = {c['id'] for c in a['candidates']}
    if ids != {c['id'] for c in view['candidates']} or len(ids) != len(view['candidates']):
        errors.append('网页候选与正式分析不一致。')
    for section in view['sections']:
        for table in section.get('tables', []):
            if any(len(row) != len(table['columns']) for row in table['rows']):
                errors.append('表格列数不一致，必须先修正字段映射。')
    return errors


def workbook(view, out):
    from research import skill_home
    template = skill_home() / 'assets/workbook-template.xlsx'
    if not template.exists():
        template = skill_home() / 'assets/approved-workbook-example.xlsx'
    book = load_workbook(template)
    if len(book.worksheets) != 8:
        raise ValueError('模板工作表数不符合已认可格式。')
    for sheet, model in zip(book, view['sheets']):
        sheet.title = model['name']
        old_filter = sheet.auto_filter.ref
        styles = {c.coordinate: copy.copy(c._style) for row in sheet for c in row if c.has_style}
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        for row in sheet:
            for cell in row:
                cell.value = None; cell.hyperlink = None; cell.comment = None
        for address, value in model['cells'].items():
            cell = sheet[address]
            cell.value = value
            if address in styles:
                cell._style = copy.copy(styles[address])
            alignment = copy.copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
        for merged in model.get('merges', []):
            sheet.merge_cells(merged)
        for address, link in model.get('links', {}).items():
            sheet[address].hyperlink = link or None
        for row, height in model.get('heights', {}).items():
            if float(height) > 409:
                raise ValueError('行高超过Excel上限，须先修版面映射。')
            sheet.row_dimensions[int(row)].height = height
        for col, width in model.get('widths', {}).items():
            sheet.column_dimensions[col].width = width
        for row in sheet.row_dimensions.values():
            row.hidden = False
        for col in sheet.column_dimensions.values():
            col.hidden = False
        last = max(int(re.search(r'\d+', k)[0]) for k in model['cells'])
        if sheet.max_row > last:
            sheet.delete_rows(last + 1, sheet.max_row - last)
        sheet.auto_filter.filterColumn = []; sheet.auto_filter.sortState = None
        sheet.auto_filter.ref = model.get('autofilter', old_filter)
        sheet.sheet_properties.filterMode = False
        sheet.sheet_state = 'visible'
        sheet.conditional_formatting._cf_rules.clear()
        sheet.print_area = f'A1:{get_column_letter(model["max_col"])}{last}'
    book.properties.title = view['title']
    book.properties.subject = view['context']
    book.properties.description = view['version']
    book.save(out)
    book.close()
    result = load_workbook(out, data_only=False)
    for s, model in zip(result, view['sheets']):
        for address, value in model['cells'].items():
            actual = s[address].value
            if not same_cell(actual, value):
                raise ValueError(f'工作簿回读不一致：{s.title}!{address}')
    result.close()


def web(view, base, out):
    e = lambda text: html.escape(str(text), quote=True)
    from research import skill_home
    css = (skill_home() / 'assets/web/report.css').read_text('utf-8')
    paragraph = lambda text: '<p>' + e(text).replace('\n', '<br>') + '</p>'
    tabs = ''.join(f'<button role="tab" data-tab="s{i}" aria-selected="{str(i==0).lower()}">{e(s["title"])}</button>' for i, s in enumerate(view['sections']))
    parts = []
    for i, section in enumerate(view['sections']):
        body = ''.join(paragraph(p) for p in section.get('paragraphs', []))
        for f in section.get('figures', []):
            src = (base / f['path']).resolve()
            if sha(src) != f['sha256']:
                raise ValueError('图版与当前输入不一致。')
            dest = out / 'figures' / src.name
            dest.parent.mkdir(exist_ok=True)
            if dest.exists() and sha(dest) != sha(src):
                raise ValueError('同名图版内容冲突。')
            shutil.copy2(src, dest)
            body += f'<figure class="graphic"><a href="figures/{e(src.name)}"><img src="figures/{e(src.name)}" alt="{e(f["label"])}" style="display:block;width:100%;height:auto"></a><figcaption>{e(f.get("caption", ""))}</figcaption></figure>'
        for j, t in enumerate(section.get('tables', [])):
            body += '<div class="rich-table"><div class="table-head"><h3>' + e(t['title']) + '</h3></div><div class="table-scroll"><table><thead><tr>'
            body += ''.join('<th>' + e(x) + '</th>' for x in t['columns']) + '</tr></thead><tbody>'
            for row in t['rows']:
                body += '<tr>' + ''.join('<td>' + e(v).replace('\n', '<br>') + '</td>' for v in row) + '</tr>'
            body += '</tbody></table></div></div>'
        if section.get('show_candidates'):
            body += '<label>查找物质 <input id="candidate-search" type="search"></label>'
            for c in view['candidates']:
                body += f'<details class="candidate more" data-search="{e(c["title"]+" "+c["group"])}"><summary>{e(c["title"])} · {e(c["group"])}</summary>'
                body += paragraph(c['summary'])
                for detail in c.get('details', []):
                    body += '<h4>' + e(detail['label']) + '</h4>' + paragraph(detail['text'])
                    for source in detail.get('sources', []):
                        if not source['url'].startswith(('https://', 'http://')):
                            raise ValueError('来源必须是明确的网页地址。')
                        body += '<a href="' + e(source['url']) + '" target="_blank" rel="noopener">' + e(source['label']) + '</a> '
                body += '</details>'
        parts.append(f'<section id="s{i}" class="page" {"hidden" if i else ""}><div class="controls"><button data-copy="s{i}">复制正文与表格</button></div><article class="reading">{body}</article></section>')
    script = '''document.querySelectorAll('[data-tab]').forEach(b=>b.addEventListener('click',()=>{document.querySelectorAll('.page').forEach(p=>p.hidden=p.id!==b.dataset.tab);document.querySelectorAll('[data-tab]').forEach(x=>x.setAttribute('aria-selected',x===b));location.hash=b.dataset.tab;}));
const chosen=document.querySelector('[data-tab="'+location.hash.slice(1)+'"]');if(chosen)chosen.click();
document.querySelectorAll('[data-copy]').forEach(b=>b.addEventListener('click',async()=>{const el=document.querySelector('#'+b.dataset.copy+' article'),copy=el.cloneNode(true);copy.querySelectorAll('details').forEach(x=>x.open=true);copy.querySelectorAll('input,label').forEach(x=>x.remove());try{await navigator.clipboard.write([new ClipboardItem({'text/html':new Blob([copy.innerHTML],{type:'text/html'}),'text/plain':new Blob([copy.textContent],{type:'text/plain'})})]);b.textContent='已复制';}catch(e){const range=document.createRange();range.selectNodeContents(el);getSelection().removeAllRanges();getSelection().addRange(range);b.textContent='已选中，请按Ctrl+C';}}));
const search=document.getElementById('candidate-search');if(search)search.addEventListener('input',()=>document.querySelectorAll('.candidate').forEach(x=>x.hidden=!x.dataset.search.toLowerCase().includes(search.value.toLowerCase())));'''
    page = '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + e(view['title']) + '</title><style>' + css + '</style><header class="masthead"><div><p id="context">' + e(view['context']) + ' · ' + e(view['version']) + '</p><h1>' + e(view['title']) + '</h1></div><a class="button" href="analysis.xlsx">分析工作簿</a></header><nav class="tabs">' + tabs + '</nav><main>' + ''.join(parts) + '</main><script>' + script + '</script></html>'
    (out / 'index.html').write_text(page, 'utf-8')


def render(path, output):
    path, output = Path(path).resolve(), Path(output).resolve()
    if output.exists():
        raise ValueError('保留旧版，请指定新的输出目录。')
    spec = load(path)
    analysis = load(path.parent / spec['analysis'])
    spec['publication'] = analysis['publication']
    errors = validate(spec, path.parent)
    if errors:
        raise ValueError('；'.join(errors))
    output.mkdir(parents=True)
    write(output / 'render_status.json', {'status': 'rendering', 'release_approved': False})
    try:
        workbook(spec['publication'], output / 'analysis.xlsx')
        web(spec['publication'], path.parent, output)
        write(output / 'render_status.json', {'status': 'rendered_for_display_review', 'release_approved': False,
              'analysis_sha256': spec['analysis_sha256'], 'publication_sha256': digest(spec['publication']),
              'new_scientific_judgment': False, 'files': {p.name: sha(p) for p in output.iterdir() if p.name in ('analysis.xlsx', 'index.html')}})
    except BaseException as exc:
        write(output / 'render_status.json', {'status': 'failed', 'release_approved': False, 'error': str(exc)})
        raise
    return {'output': str(output), 'status': 'rendered_for_display_review', 'release_approved': False}


if __name__ == '__main__':
    import argparse
    p = argparse.ArgumentParser(); p.add_argument('presentation', type=Path); p.add_argument('--output', type=Path, required=True)
    args = p.parse_args()
    try:
        print(json.dumps(render(args.presentation, args.output), ensure_ascii=False))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
