"""Local, disk-backed exports. Catalog input remains read-only."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import queue
import sqlite3
import threading
import time
import uuid
import zipfile
from pathlib import Path


class ExportCancelled(Exception):
    pass


def safe_cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


class ExportJobs:
    """One local export worker; repeated requests reuse their existing job."""

    def __init__(self, api, output_root, catalog_path):
        self.api = api
        self.root = Path(output_root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.catalog_path = Path(catalog_path).resolve()
        self.lock = threading.RLock()
        self.jobs = {}
        self.cancelled = set()
        self.pending = queue.Queue()
        self.support = {}
        for manifest in self.root.glob('*/job.json'):
            try:
                job = json.loads(manifest.read_text(encoding='utf-8'))
                if job['status'] in ('queued', 'running'):
                    job.update(status='error', error='服务已重启，请重新生成。已写入的临时文件保留在本机。')
                self.jobs[job['id']] = job
                self._save(job)
            except (OSError, ValueError, KeyError):
                continue
        threading.Thread(target=self._worker, daemon=True, name='catalog-export').start()

    def _save(self, job):
        folder = self.root / job['id']
        folder.mkdir(exist_ok=True)
        target = folder / 'job.json'
        temporary = folder / 'job.json.tmp'
        temporary.write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding='utf-8')
        temporary.replace(target)

    def _save_best_effort(self, job):
        try:
            self._save(job)
        except OSError as error:
            job['persistence_error'] = str(error)

    def create(self, args, format):
        if format not in ('csv', 'xlsx', 'sql'):
            raise ValueError('Unsupported export format')
        args = dict(args or {})
        for key in ('limit', 'offset', 'scope', 'page', 'page_size'):
            args.pop(key, None)
        args['take'] = 'all'
        signature = hashlib.sha256(json.dumps([str(self.catalog_path), format, args], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        with self.lock:
            for old in self.jobs.values():
                if old.get('signature') == signature and old['status'] in ('queued', 'running', 'complete'):
                    if old['status'] != 'complete' or (self.root / old['id'] / old['filename']).is_file():
                        return dict(old)
            job = dict(id=uuid.uuid4().hex, signature=signature, format=format, args=args,
                       filename='catalog-selection.' + format, status='queued', records=0,
                       bytes=None, error='', created_at=time.time())
            self.jobs[job['id']] = job
            self._save(job)
            self.pending.put(job['id'])
            return dict(job)

    def get(self, job_id):
        with self.lock:
            if job_id not in self.jobs:
                raise KeyError('Unknown export job')
            return dict(self.jobs[job_id])

    def cancel(self, job_id):
        with self.lock:
            job = self.jobs[job_id]
            if job['status'] in ('queued', 'running'):
                self.cancelled.add(job_id)
                job['cancel_requested'] = True
                if job['status'] == 'queued':
                    job['status'] = 'cancelled'
                self._save(job)
            return dict(job)

    def file(self, job_id):
        if job_id.startswith('support-'):
            if not self.support:
                self.metadata()
            if job_id not in self.support:
                raise KeyError('Unknown supporting file')
            return self.support[job_id]
        job = self.get(job_id)
        if job['status'] != 'complete':
            raise ValueError('The export is not ready to download')
        path = (self.root / job_id / job['filename']).resolve()
        if not path.is_relative_to(self.root) or not path.is_file():
            raise ValueError('The export file is unavailable')
        return path

    def _check_cancel(self, job):
        if job['id'] in self.cancelled:
            raise ExportCancelled()

    def _rows(self, job):
        for part in self.api.iter_batches(job['args']):
            self._check_cancel(job)
            for row in part['batch']['matches']:
                self._check_cancel(job)
                yield part['kind'], row
                job['records'] += 1
            with self.lock:
                self._save(job)

    def _worker(self):
        while True:
            job_id = self.pending.get()
            job = self.jobs[job_id]
            try:
                self._check_cancel(job)
                with self.lock:
                    job.update(status='running', started_at=time.time())
                    self._save(job)
                partial = self.root / job_id / (job['filename'] + '.partial')
                getattr(self, '_write_' + job['format'])(job, partial)
                self._check_cancel(job)
                self._verify(job, partial)
                self._check_cancel(job)
                final = partial.with_suffix('')
                partial.replace(final)
                with self.lock:
                    job.update(status='complete', bytes=final.stat().st_size, finished_at=time.time(),
                               download_url='/api/workspace/export-file?id=' + job_id)
                    self._save(job)
            except ExportCancelled:
                with self.lock:
                    job.update(status='cancelled', finished_at=time.time())
                    self._save_best_effort(job)
            except Exception as error:
                with self.lock:
                    job.update(status='error', error=str(error), finished_at=time.time())
                    self._save_best_effort(job)
            finally:
                self.pending.task_done()

    def _write_csv(self, job, path):
        fields = ['collection', 'record_id', 'name', 'value', 'unit', 'species', 'support_type', 'record_json']
        with path.open('w', encoding='utf-8-sig', newline='') as output:
            writer = csv.DictWriter(output, fields)
            writer.writeheader()
            for kind, row in self._rows(job):
                writer.writerow(dict(collection=kind,
                    record_id=safe_cell(row.get('record_id') or row.get('source_id') or row.get('gene_id') or row.get('protein_id') or row.get('pathway_id')),
                    name=safe_cell(row.get('name') or row.get('preferred_name') or row.get('property_name')),
                    value=safe_cell(row.get('value') if row.get('value') is not None else row.get('value_text')),
                    unit=safe_cell(row.get('unit')), species=safe_cell(row.get('species') or row.get('scientific_name')),
                    support_type=row.get('support_type'), record_json=json.dumps(row, ensure_ascii=False)))

    def _write_sql(self, job, path):
        with path.open('w', encoding='utf-8', newline='\n') as output:
            output.write('BEGIN;\nCREATE TABLE IF NOT EXISTS catalog_records(collection TEXT, record_json TEXT);\n')
            for kind, row in self._rows(job):
                payload = json.dumps(row, ensure_ascii=False).replace("'", "''")
                output.write("INSERT INTO catalog_records VALUES ('" + kind.replace("'", "''") + "','" + payload + "');\n")
            output.write('COMMIT;\n')

    def _write_xlsx(self, job, path):
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.worksheet._writer import WorksheetWriter
        class LocalWorksheetWriter(WorksheetWriter):
            def cleanup(self):
                Path(self.out).unlink(missing_ok=True)
        workbook = Workbook(write_only=True)
        fields = ['record_id', 'source_id', 'name', 'chemical_name', 'property_name', 'value', 'value_text', 'unit', 'species', 'scientific_name', 'support_type', 'source_url', 'record_json']
        sheets = {}
        temporary = path.parent / 'worksheet-parts'
        temporary.mkdir(exist_ok=True)

        def append(kind, headers, values):
            state = sheets.setdefault(kind, [None, 0, 0])
            if state[0] is None or state[1] >= 1048575:
                state[2] += 1
                sheet = workbook.create_sheet((kind[:24] + '_' + str(state[2]))[:31])
                # Keep worksheet staging on the same local data disk, without changing global temp settings.
                sheet._writer = LocalWorksheetWriter(sheet, out=str(temporary / (sheet.title + '.xml')))
                sheet.append(headers)
                state[0], state[1] = sheet, 0
            cells = []
            for value in values:
                if isinstance(value, str):
                    cell = WriteOnlyCell(state[0], value=value)
                    cell.data_type = 's'
                    cells.append(cell)
                else:
                    cells.append(value)
            state[0].append(cells)
            state[1] += 1

        try:
            for kind, row in self._rows(job):
                values = []
                for field in fields:
                    value = json.dumps(row, ensure_ascii=False) if field == 'record_json' else row.get(field)
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False)
                    if isinstance(value, str) and len(value) > 32000:
                        key = row.get('record_id') or row.get('source_id') or row.get('gene_id') or row.get('protein_id')
                        for start in range(0, len(value), 32000):
                            append('long_fields', ['record_id', 'field', 'part', 'text'], [key, field, start // 32000 + 1, value[start:start + 32000]])
                        value = 'See long_fields: ' + str(key) + ' / ' + field
                    values.append(value)
                append(kind, fields, values)
            if not sheets:
                append('records', fields, [])
            workbook.save(path)
        except Exception:
            for sheet in workbook.worksheets:
                try:
                    sheet.close()
                except Exception:
                    pass
            raise

    def _verify(self, job, path):
        if not path.is_file() or not path.stat().st_size:
            raise ValueError('Export produced no file')
        if job['format'] == 'xlsx':
            with zipfile.ZipFile(path) as archive:
                bad = archive.testzip()
                if bad:
                    raise ValueError('XLSX CRC verification failed: ' + bad)
            from openpyxl import load_workbook
            with path.open('rb') as stream:
                workbook = load_workbook(stream, read_only=True)
                try:
                    for sheet in workbook:
                        next(sheet.iter_rows(values_only=True), None)
                finally:
                    workbook.close()
        elif job['format'] == 'csv':
            with path.open(encoding='utf-8-sig', newline='') as stream:
                reader = csv.DictReader(stream)
                first = next(reader, None)
                if first:
                    json.loads(first['record_json'])
        else:
            # Restore a small real sample; the complete output retains its transaction footer.
            connection = sqlite3.connect(':memory:')
            try:
                with path.open(encoding='utf-8') as stream:
                    for number, line in enumerate(stream):
                        connection.execute(line)
                        if number >= 5:
                            break
                for (payload,) in connection.execute('SELECT record_json FROM catalog_records LIMIT 3'):
                    json.loads(payload)
                with path.open('rb') as stream:
                    stream.seek(max(0, path.stat().st_size - 32))
                    if not stream.read().rstrip().endswith(b'COMMIT;'):
                        raise ValueError('SQL export is incomplete')
            finally:
                connection.close()
        job['verification'] = 'file_readback_passed'

    def metadata(self):
        """Create small support artifacts once; their bytes and downloads share the same files."""
        with self.lock:
            if not self.support:
                from catalog_queries import Catalog, ro
                directory = self.root / 'supporting'
                directory.mkdir(exist_ok=True)
                dictionary = io.StringIO(newline='')
                writer = csv.writer(dictionary)
                writer.writerow(['store', 'table', 'field', 'type', 'required', 'primary_key'])
                schema = io.StringIO()
                with Catalog(self.catalog_path) as catalog:
                    stores = {key: value for key, value in catalog.paths.items() if value.suffix == '.sqlite3'}
                for name, path in stores.items():
                    connection = ro(path)
                    try:
                        for table, sql, kind in connection.execute('SELECT name,sql,type FROM sqlite_master WHERE sql IS NOT NULL ORDER BY type,name'):
                            if kind == 'table':
                                for field in connection.execute('PRAGMA table_info("' + table.replace('"', '""') + '")'):
                                    writer.writerow([name, table, field[1], field[2], field[3], field[5]])
                            if name == 'fish_database':
                                schema.write('-- ' + name + '\n' + sql + ';\n\n')
                    finally:
                        connection.close()
                contents = {
                    'dictionary': ('data-dictionary.csv', ('\ufeff' + dictionary.getvalue()).encode()),
                    'schema': ('schema.sql', ('\ufeff' + schema.getvalue()).encode()),
                    'manifest': ('catalog-files.json', json.dumps(self.api.files(), ensure_ascii=False, indent=2).encode()),
                }
                for key, (filename, body) in contents.items():
                    path = directory / filename
                    path.write_bytes(body)
                    self.support['support-' + key] = path
            return {key.removeprefix('support-'): dict(filename=path.name, bytes=path.stat().st_size,
                    path='/api/workspace/export-file?id=' + key) for key, path in self.support.items()}
