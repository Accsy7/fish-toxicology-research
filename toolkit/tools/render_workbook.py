"""Template row/block renderer. Profiles choose sheets, fields, widths and styles."""
import copy
import math
import unicodedata
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "vendor/python"))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from report_model import field, mapped, select, text


def text_units(value):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value))


def line_count(value, width):
    return sum(max(1, math.ceil(text_units(line) / max(2, width - 3)))
               for line in text(value).split("\n"))


def write_cell(ws, r, c, value, sample, expected):
    dest = ws.cell(r, c)
    # Copy style objects: style IDs are local to each workbook's registry.
    for attribute in ("font", "fill", "border", "protection", "alignment"):
        setattr(dest, attribute, copy.copy(getattr(sample, attribute)))
    dest.number_format = sample.number_format
    if isinstance(value, dict):
        link = value.get("url")
        value = value.get("text", "")
        if link:
            if link.startswith("#"):
                dest.hyperlink = Hyperlink(ref=dest.coordinate, location=link[1:])
            elif link.startswith(("https://", "http://")):
                dest.hyperlink = link
            else:
                raise ValueError(f"Unsafe hyperlink: {link}")
            font = copy.copy(dest.font)
            font.color, font.underline = "285E78", "single"
            dest.font = font
    if isinstance(value, (list, dict)):
        value = text(value)
    if isinstance(value, str) and value.startswith(("https://", "http://")) and "\n" not in value:
        dest.hyperlink = value
        font = copy.copy(dest.font)
        font.color, font.underline = "285E78", "single"
        dest.font = font
    dest.value = value
    # Source strings beginning with '=' are text; formulas need an explicit profile operation.
    if isinstance(value, str):
        dest.data_type = "s"
    dest.alignment = Alignment(horizontal="right" if type(value) in (int, float) else "left",
                               vertical="top", wrap_text=True)
    expected.append((ws.title, dest.coordinate, value))
    return dest


def fit_row(ws, row, spans, minimum=23, options=None):
    options = options or {}
    height = minimum
    for start, end in spans:
        cell = ws.cell(row, start)
        width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(start, end + 1))
        height = max(height, line_count(cell.value, width) * (cell.font.sz or 11) * options.get('line_spacing', 1.55) + options.get('padding', 8))
    # Fail visibly instead of silently clipping very long text.
    if height > 409:
        raise ValueError(f"{ws.title}!{row}: text exceeds one Excel row; split the record into detail rows")
    ws.row_dimensions[row].height = height


def split_text(value, width, font_size, options):
    """Keep every character; long detail blocks continue in additional template rows."""
    value = text(value)
    limit = min(409, options.get('maximum', 380))
    spacing, padding = options.get('line_spacing', 1.55), options.get('padding', 8)
    if limit < font_size * spacing + padding:
        raise ValueError('row_height.maximum must fit at least one line')
    if not value:
        return ['']
    parts = []
    while value:
        lo, hi = 1, len(value)
        while lo <= hi:
            mid = (lo + hi) // 2
            if line_count(value[:mid], width) * font_size * spacing + padding <= limit:
                lo = mid + 1
            else:
                hi = mid - 1
        if hi == 0:
            raise ValueError('Text cannot fit the configured detail width')
        cut = hi
        if hi < len(value):
            boundaries = [value.rfind(mark, max(0, hi // 2), hi) for mark in ('\n', '。', '；', '. ')]
            boundary = max(boundaries)
            if boundary >= 0:
                cut = boundary + 1
        parts.append(value[:cut])
        value = value[cut:]
    return parts


def render(data, profile, template, output):
    source = load_workbook(template)
    book = Workbook()
    book.remove(book.active)
    expected, anchors = [], {}
    for cfg in profile["workbook"]["sheets"]:
        sample = source.worksheets[cfg["template_sheet"]] if isinstance(cfg["template_sheet"], int) else source[cfg["template_sheet"]]
        style_ref = cfg.get("table_style_sheet", cfg["template_sheet"])
        table_sample = source.worksheets[style_ref] if isinstance(style_ref, int) else source[style_ref]
        ws = book.create_sheet(cfg["name"])
        for key, dim in sample.column_dimensions.items():
            dest = ws.column_dimensions[key]
            dest.width, dest.outlineLevel = dim.width, dim.outlineLevel
            if dim.has_style:
                for attribute in ("font", "fill", "border", "protection", "alignment"):
                    setattr(dest, attribute, copy.copy(getattr(dim, attribute)))
                dest.number_format = dim.number_format
        for key, width in cfg.get("widths", {}).items():
            ws.column_dimensions[key].width = width
        ws.sheet_view.showGridLines = sample.sheet_view.showGridLines
        ws.sheet_view.zoomScale = cfg.get("zoom", 85)
        ws.sheet_properties = copy.copy(sample.sheet_properties)
        ws.sheet_properties.filterMode = False
        ws.sheet_format = copy.copy(sample.sheet_format)
        ws.page_setup = copy.copy(sample.page_setup)
        ws.page_margins = copy.copy(sample.page_margins)
        ws.freeze_panes = cfg.get("freeze", "A4")
        ws.sheet_state = "visible"
        ncols = cfg["columns_count"]

        def line(r, value, sr, c=1, end=None):
            end = end or ncols
            write_cell(ws, r, c, value, sample.cell(sr, c), expected)
            if end > c:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=end)
            fit_row(ws, r, [(c, end)], 26 if sr == 1 else 23)

        line(1, cfg.get("title", cfg["name"]), 1)
        line(2, data.get("context", ""), 2)
        if cfg["kind"] == "blocks":
            r = 3
            for item in select(data, cfg["collection"]):
                anchors[item["id"]] = f"'{ws.title}'!A{r}"
                line(r, f"{item['name']} · {item.get('group', '')}", cfg.get("block_title_sample", 3))
                r += 1
                rows = [{"label": "预测", "text": item.get("prediction", "")}] + item.get("details", [])
                for detail in rows:
                    body = detail.get("text", "")
                    if detail.get("source_ids"):
                        body = text(body) + "\n来源：" + " / ".join(detail["source_ids"])
                    start = cfg['label_end'] + 1
                    width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(start, ncols + 1))
                    options = cfg.get('row_height', {})
                    for i, part in enumerate(split_text(body, width, sample['C4'].font.sz or 11, options)):
                        label = detail['label'] + (cfg.get('continuation_label', '（续）') if i else '')
                        write_cell(ws, r, 1, label, sample['A4'], expected)
                        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cfg['label_end'])
                        write_cell(ws, r, start, part, sample['C4'], expected)
                        ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=ncols)
                        fit_row(ws, r, [(1, cfg['label_end']), (start, ncols)], options=options)
                        r += 1
                r += 1
        else:
            r = cfg.get("header_row", 3)
            for j, col in enumerate(cfg["columns"], 1):
                write_cell(ws, r, j, col["label"], table_sample.cell(cfg["header_sample"], j), expected)
            fit_row(ws, r, [(j, j) for j in range(1, ncols+1)], 24)
            header = r
            for item in select(data, cfg["collection"]):
                r += 1
                for j, col in enumerate(cfg["columns"], 1):
                    value = mapped(item, col)
                    if col.get('detail_link') and not value:
                        value = col.get('detail_label', '查看')
                    write_cell(ws, r, j, value, table_sample.cell(cfg["body_sample"], j), expected)
                fit_row(ws, r, [(j, j) for j in range(1, ncols+1)])
            if r == header and cfg.get("empty"):
                r += 1
                line(r, cfg["empty"], 2)
            ws.auto_filter.ref = f"A{header}:{get_column_letter(ncols)}{r}"
            r += 1
        # Optional supplemental tables use the same template styles.
        for extra in cfg.get("append", []):
            rows = select(data, extra["collection"])
            if not rows:
                continue
            r += 1
            line(r, extra["title"], 1)
            r += 1
            for j, col in enumerate(extra["columns"], 1):
                write_cell(ws, r, j, col["label"], sample.cell(cfg.get("header_sample", 3), j), expected)
            fit_row(ws, r, [(j, j) for j in range(1, len(extra["columns"])+1)])
            for item in rows:
                r += 1
                for j, col in enumerate(extra["columns"], 1):
                    write_cell(ws, r, j, mapped(item, col), sample.cell(cfg.get("body_sample", 4), j), expected)
                fit_row(ws, r, [(j, j) for j in range(1, len(extra["columns"])+1)])
            r += 1
        ws.print_area = f"A1:{get_column_letter(ncols)}{max(2,r-1)}"
        ws.print_options.horizontalCentered = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"1:{cfg.get('header_row', 2 if cfg['kind']=='blocks' else 3)}"
    # Resolve links after all detail blocks have been laid out.
    for cfg, ws in zip(profile["workbook"]["sheets"], book):
        if cfg["kind"] != "table":
            continue
        for j, col in enumerate(cfg["columns"], 1):
            if col.get("detail_link"):
                for k, item in enumerate(select(data, cfg["collection"]), cfg.get("header_row", 3)+1):
                    if item["id"] in anchors:
                        ws.cell(k, j).hyperlink = Hyperlink(ref=ws.cell(k,j).coordinate, location=anchors[item["id"]])
                        font = copy.copy(ws.cell(k,j).font)
                        font.color, font.underline = "285E78", "single"
                        ws.cell(k,j).font = font
    book.properties.title, book.properties.subject = data["title"], data.get("context", "")
    book.save(output)
    book.close()
    source.close()
    actual = load_workbook(output)
    for name, address, value in expected:
        got = actual[name][address].value
        equal = math.isclose(got, value, rel_tol=1e-14, abs_tol=0) if type(got) in (int,float) and type(value) in (int,float) else got == value or (got is None and value == "")
        if not equal:
            raise ValueError(f"Readback differs: {name}!{address}")
    for ws in actual:
        if any(d.hidden for d in ws.row_dimensions.values()) or any(d.hidden for d in ws.column_dimensions.values()):
            raise ValueError(f"Hidden data in {ws.title}")
    actual.close()
    return {"checked_cells": len(expected), "sheets": len(profile["workbook"]["sheets"]), "detail_anchors": anchors}
