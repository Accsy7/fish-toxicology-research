"""Serve the preserved frontend with local, read-only catalog queries."""
from pathlib import Path
import sys, argparse, json, csv, io, mimetypes, threading, tempfile, os, sqlite3, queue, time, select, socket
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from urllib.parse import urlparse,parse_qs,unquote
ROOT=Path(__file__).resolve().parent
for p in [ROOT/'runtime',ROOT/'../runtime',ROOT/'../vendor/catalog-runtime',ROOT/'vendor/catalog-runtime',ROOT/'../../chemical_expansion_20261018/runtime',ROOT/'../chemical_expansion_20261018/runtime']:
    if p.exists():sys.path.insert(0,str(p.resolve()))
# The local revision contains the current query adapters. Dependency runtimes
# supply packages and fallback modules, but must not shadow these maintained files.
sys.path=[entry for entry in sys.path if entry!=str(ROOT)]
sys.path.insert(0,str(ROOT))
from catalog_queries import Catalog, normalize_query
from extended_api import WorkspaceAPI

def safe_cell(v):
    if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False)
    if isinstance(v,str) and v.startswith(('=','+','-','@')):return "'"+v
    return v

class Handler(BaseHTTPRequestHandler):
    _export_lock = threading.Lock()
    def export_jobs(self):
        from export_jobs import ExportJobs
        with self._export_lock:
            if not hasattr(self.server, 'exports'):
                self.server.exports = ExportJobs(self.server.api, ROOT/'export-cache', self.server.catalog)
        return self.server.exports

    def do_POST(self):
        origin = self.headers.get('Origin')
        if origin and origin != 'http://' + self.headers.get('Host', ''):
            return self.send(403, b'{"error":"Local origin required"}')
        try:
            length = int(self.headers.get('Content-Length', 0))
            if length <= 0: raise ValueError('A JSON request body is required')
            body = json.loads(self.rfile.read(length))
            if not isinstance(body, dict):raise ValueError('Query arguments must be a JSON object')
            if urlparse(self.path).path == '/api/workspace/export-jobs':
                return self.json(self.export_jobs().create(body.get('args', {}), body.get('format', 'csv')))
            if urlparse(self.path).path == '/api/workspace/export-jobs/cancel':
                return self.json(self.export_jobs().cancel(body.get('id', '')))
            if self.read_query(urlparse(self.path).path, body):return
            return self.send(404, b'{"error":"Unknown endpoint"}')
        except (BrokenPipeError,ConnectionResetError):pass
        except (ValueError, TypeError, KeyError, OSError) as e:
            return self.send(400, json.dumps({'error':str(e)}, ensure_ascii=False).encode())
        except Exception as e:
            return self.send(500, json.dumps({'error':str(e)}, ensure_ascii=False).encode())
    def send(self,status,body,ctype='application/json; charset=utf-8',download=None):
        self.send_response(status);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(body)))
        self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff')
        if download:self.send_header('Content-Disposition','attachment; filename="'+download+'"')
        self.end_headers();self.wfile.write(body)
    def json(self,data):self.send(200,json.dumps(data,ensure_ascii=False,allow_nan=False).encode())

    def cancellable_json(self,call):
        """Interrupt only this request's SQL when its client closes the socket."""
        from query_runtime import QueryRuntime
        runtime=QueryRuntime();finished=threading.Event()

        def watch_connection():
            while not finished.is_set():
                try:
                    readable,_,failed=select.select([self.connection],[],[self.connection],.2)
                    if finished.is_set():return
                    if failed or (readable and not self.connection.recv(1,socket.MSG_PEEK)):
                        runtime.cancel();return
                    # Buffered/pipelined input is not a disconnect. Do not consume
                    # it or spin while this handler still owns the current query.
                    if readable:finished.wait(.2)
                except OSError:
                    if not finished.is_set():runtime.cancel()
                    return

        watcher=threading.Thread(target=watch_connection,daemon=True,name='catalog-query-watch')
        watcher.start()
        try:
            with runtime:
                runtime.check();result=call();runtime.check()
                self.json(result)
        except Exception:
            if not runtime.stopped.is_set():raise
            self.close_connection=True
        finally:
            finished.set();watcher.join(timeout=.5)

    def send_error(self,code,message=None,explain=None):
        # BaseHTTPRequestHandler rejects >64 KiB request lines before do_GET.
        # Its default HTML error caused response.json() to fail on '<!DOCTYPE'.
        # Keep protocol errors machine-readable, including that legacy case.
        self.close_connection=True
        text=message or self.responses.get(code,('Request error',))[0]
        if code==414:text='Query URL is too long; send the same filters as a JSON POST body.'
        body=json.dumps({'error':text,'status':code},ensure_ascii=False).encode('utf-8')
        self.send_response(code);self.send_header('Content-Type','application/json; charset=utf-8')
        self.send_header('Content-Length',str(len(body)));self.send_header('Connection','close')
        self.send_header('Cache-Control','no-store');self.end_headers()
        if getattr(self,'command',None)!='HEAD':self.wfile.write(body)

    def read_query(self,path,args):
        """GET and POST share one read-only query implementation."""
        if path=='/api/workspace/search':
            self.cancellable_json(lambda:self.server.api.search(args));return True
        calls={
            '/api/workspace/meta':lambda:self.server.api.meta(),
            '/api/workspace/detail':lambda:self.server.api.detail(args),
            '/api/workspace/network':lambda:self.server.api.network(args),
            '/api/workspace/profile':lambda:self.server.api.legacy_context(args),
            '/api/workspace/reader':lambda:self.server.api.reader(args),
            '/api/workspace/source-material':lambda:self.server.api.source_material(args),
            '/api/workspace/files':lambda:self.server.api.files(),
            '/api/workspace/export-jobs':lambda:self.export_jobs().get(args.get('id','')),
            '/api/workspace/download-metadata':lambda:self.export_jobs().metadata(),
        }
        if path in calls:self.json(calls[path]());return True
        if path=='/api/workspace/stream':self.stream(args);return True
        if path=='/api/query':
            with Catalog(self.server.catalog) as catalog:
                result=catalog.query(args.get('kind','coverage'),args.get('term',''),args.get('species',''),
                                     args.get('support','default'),int(args.get('limit',50)),
                                     int(args.get('offset',0)),args.get('assertion','any'),
                                     sort_field=args.get('sort_field'),sort_direction=args.get('sort_direction','asc'))
            self.json(result);return True
        return False

    def do_GET(self):
        u=urlparse(self.path);a={k:v[0] for k,v in parse_qs(u.query,keep_blank_values=True).items()}
        try:
            if self.read_query(u.path,a):return
            if u.path=='/api/workspace/export-file':return self.file(self.export_jobs().file(a.get('id','')))
            if u.path in ('/api/workspace/schema.sql','/api/workspace/dictionary.csv'):
                return self.file(self.export_jobs().file('support-dictionary' if u.path.endswith('.csv') else 'support-schema'))
            if u.path=='/api/workspace/file':return self.file(self.server.api.file_path(a.get('id','')))
            if u.path=='/api/workspace/export.xlsx':return self.xlsx(a)
            if u.path=='/api/workspace/export.sql':return self.sql(a)
            if u.path=='/api/workspace/export.csv':return self.export(a)
            if u.path=='/api/workspace/detail.json':
                return self.send(200,json.dumps(self.server.api.detail(a),ensure_ascii=False,indent=2).encode(),download='source-record.json')
            if u.path in ('/api/query','/api/structure.mol','/api/sequence.txt','/api/pathway.kgml'):
                kind={'/api/structure.mol':'structure','/api/sequence.txt':'sequence','/api/pathway.kgml':'graph'}.get(u.path,a.get('kind','coverage'))
                with Catalog(self.server.catalog) as c:r=c.query(kind,a.get('term',''),a.get('species',''),a.get('support','default'),int(a.get('limit',50)),int(a.get('offset',0)),a.get('assertion','any'))
                if u.path=='/api/query':return self.json(r)
                if kind=='structure':return self.send(200,r['molfile'].encode(),'chemical/x-mdl-molfile',download='structure-'+a.get('term','')+'.mol')
                if kind=='graph':return self.send(200,r['xml'].encode(),'application/xml',download='pathway.kgml')
                if kind=='sequence':
                    sequence=r['sequence'];fasta='>'+r['accession']+' '+r['species']+'\n'+'\n'.join(sequence[n:n+80] for n in range(0,len(sequence),80))+'\n'
                    return self.send(200,fasta.encode(),'text/plain; charset=utf-8',download='sequence.fasta')
                return self.json(r)
            if u.path=='/api/structure.svg':
                from rdkit import Chem
                from rdkit.Chem.Draw import rdMolDraw2D
                with Catalog(self.server.catalog) as c:r=c.query('structure',a.get('term',''))
                mol=Chem.MolFromMolBlock(r['molfile'],removeHs=False)
                if mol is None:raise ValueError('The source MOL could not be rendered')
                drawer=rdMolDraw2D.MolDraw2DSVG(460,280);drawer.DrawMolecule(mol);drawer.FinishDrawing()
                return self.send(200,drawer.GetDrawingText().encode(),'image/svg+xml')
            if u.path.startswith('/api/'):return self.send(404,b'{"error":"Unknown endpoint"}')
            rel=unquote(u.path).lstrip('/') or 'index.html';target=(self.server.static/rel).resolve()
            if not target.is_relative_to(self.server.static) or not target.is_file():return self.send(404,b'Not found','text/plain')
            ctype=mimetypes.guess_type(target.name)[0] or 'application/octet-stream'
            if ctype.startswith('text/'):ctype+='; charset=utf-8'
            return self.send(200,target.read_bytes(),ctype)
        except (BrokenPipeError,ConnectionResetError):pass
        except (ValueError,TypeError,KeyError,OSError) as e:self.send(400,json.dumps({'error':str(e)},ensure_ascii=False).encode())
        except Exception as e:self.send(500,json.dumps({'error':str(e)},ensure_ascii=False).encode())
    def start_stream(self,ctype,filename=None):
        self.send_response(200);self.send_header('Content-Type',ctype)
        if filename:self.send_header('Content-Disposition','attachment; filename="'+filename+'"')
        self.send_header('Cache-Control','no-store');self.send_header('Connection','close');self.end_headers();self.close_connection=True

    def stream(self,a):
        self.start_stream('application/x-ndjson; charset=utf-8')
        from query_runtime import QueryRuntime
        runtime=QueryRuntime();messages=queue.Queue(maxsize=2);started=time.monotonic()
        finished=object();state={'phase':'querying','loaded':0,'related_loaded':0}

        def enqueue(value):
            while not runtime.stopped.is_set():
                try:messages.put(value,timeout=.2);return
                except queue.Full:pass

        def produce():
            try:
                with runtime:
                    for part in self.server.api.iter_batches(a,network=a.get('mode')=='network'):
                        runtime.check();enqueue(part)
            except Exception as error:
                if not runtime.stopped.is_set():enqueue({'error':str(error)})
            finally:enqueue(finished)

        def write(part):
            self.wfile.write((json.dumps(part,ensure_ascii=False,allow_nan=False)+'\n').encode());self.wfile.flush()

        worker=threading.Thread(target=produce,daemon=True,name='catalog-stream')
        try:
            write({'type':'progress','phase':'accepted','loaded':0,'elapsed_seconds':0})
            worker.start()
            while True:
                try:part=messages.get(timeout=1)
                except queue.Empty:
                    # A real elapsed-time heartbeat also detects a closed client
                    # while its SQL statement is still executing.
                    write({'type':'progress',**state,'elapsed_seconds':round(time.monotonic()-started,2)})
                    continue
                if part is finished:
                    write({'complete':True,'elapsed_seconds':round(time.monotonic()-started,3)})
                    break
                if part.get('error'):
                    write(part);break
                if part.get('type')=='progress':state.update({k:v for k,v in part.items() if k!='type'})
                else:state.update(phase='streaming',kind=part.get('kind'),loaded=part.get('loaded',0),related_loaded=part.get('related_loaded',0))
                write({**part,'elapsed_seconds':round(time.monotonic()-started,3)})
        except (BrokenPipeError,ConnectionResetError,ConnectionAbortedError):pass
        except Exception as error:
            try:write({'error':str(error)})
            except (BrokenPipeError,ConnectionResetError,ConnectionAbortedError):pass
        finally:
            runtime.cancel()
            if worker.is_alive():worker.join(timeout=2)

    def file(self,path):
        self.send_response(200);self.send_header('Content-Type',mimetypes.guess_type(path.name)[0] or 'application/octet-stream')
        self.send_header('Content-Length',str(path.stat().st_size))
        from urllib.parse import quote
        self.send_header('Content-Disposition',"attachment; filename*=UTF-8''"+quote(path.name));self.end_headers()
        with path.open('rb') as f:
            for block in iter(lambda:f.read(1024*1024),b''):self.wfile.write(block)

    def export(self,a):
        if a.get('kinds') or a.get('take'):
            self.start_stream('text/csv; charset=utf-8','catalog-selection.csv');self.wfile.write(b'\xef\xbb\xbf')
            fields=['collection','record_id','name','value','unit','species','support_type','record_json']
            header=True
            for part in self.server.api.iter_batches(a):
                s=io.StringIO(newline='');w=csv.DictWriter(s,fields)
                if header:w.writeheader();header=False
                for r in part['batch']['matches']:
                    w.writerow({'collection':part['kind'],'record_id':safe_cell(r.get('record_id') or r.get('source_id') or r.get('gene_id') or r.get('protein_id') or r.get('pathway_id')),'name':safe_cell(r.get('name') or r.get('preferred_name') or r.get('property_name')),'value':safe_cell(r.get('value') if r.get('value') is not None else r.get('value_text')),'unit':safe_cell(r.get('unit')),'species':safe_cell(r.get('species') or r.get('scientific_name')),'support_type':r.get('support_type'),'record_json':safe_cell(json.dumps(r,ensure_ascii=False))})
                self.wfile.write(s.getvalue().encode());self.wfile.flush()
            return
        # Use exactly the same predicates and deterministic ordering as Search.
        a=dict(a);all_rows=a.pop('scope','page')=='all';a['limit']=500 if all_rows else int(a.get('limit',50))
        if all_rows:a['offset']=0
        first=self.server.api.search(a);records=first.get('matches',[])
        fields=list(dict.fromkeys(k for r in records for k in r))+['record_json']
        self.send_response(200);self.send_header('Content-Type','text/csv; charset=utf-8');self.send_header('Content-Disposition','attachment; filename="'+a.get('kind','records')+'-'+('filtered' if all_rows else 'page')+'.csv"');self.send_header('Cache-Control','no-store');self.send_header('Connection','close');self.end_headers();self.close_connection=True
        self.wfile.write(b'\xef\xbb\xbf')
        def block(rs,header=False):
            s=io.StringIO(newline='');w=csv.DictWriter(s,fields,extrasaction='ignore')
            if header:w.writeheader()
            for row in rs:w.writerow({**{k:safe_cell(v) for k,v in row.items()},'record_json':json.dumps(row,ensure_ascii=False)})
            self.wfile.write(s.getvalue().encode());self.wfile.flush()
        if fields:block(records,True)
        while all_rows and first.get('has_more'):
            a['offset']=first['offset']+first['returned'];first=self.server.api.search(a);block(first.get('matches',[]))
            if not first.get('returned'):break

    def xlsx(self,a):
        from openpyxl import Workbook
        workbook=Workbook(write_only=True);sheets={};overflow=None
        fields=['record_id','source_id','name','chemical_name','property_name','value','value_text','unit','species','scientific_name','support_type','source_url','record_json']
        for part in self.server.api.iter_batches(a):
            kind=part['kind']
            if kind not in sheets:sheets[kind]=[None,0,0]
            state=sheets[kind]
            for r in part['batch']['matches']:
                if state[0] is None or state[1]>=1048575:
                    state[2]+=1;state[0]=workbook.create_sheet(kind+'_'+str(state[2]));state[0].append(fields);state[1]=0
                values=[]
                for field in fields:
                    v=json.dumps(r,ensure_ascii=False) if field=='record_json' else r.get(field)
                    v=safe_cell(v)
                    if isinstance(v,str) and len(v)>32000:
                        if overflow is None:overflow=workbook.create_sheet('long_fields');overflow.append(['record_id','field','part','text'])
                        key=r.get('record_id') or r.get('source_id') or r.get('gene_id') or r.get('protein_id')
                        for i in range(0,len(v),32000):overflow.append([key,field,i//32000+1,v[i:i+32000]])
                        v='See long_fields: '+str(key)+' / '+field
                    values.append(v)
                state[0].append(values);state[1]+=1
        if not sheets:workbook.create_sheet('records').append(fields)
        fd,name=tempfile.mkstemp(suffix='.xlsx');os.close(fd)
        try:
            workbook.save(name);self.file(Path(name))
        finally:Path(name).unlink(missing_ok=True)

    def sql(self,a):
        self.start_stream('application/sql; charset=utf-8','catalog-selection.sql')
        self.wfile.write(b'BEGIN;\nCREATE TABLE IF NOT EXISTS catalog_records(collection TEXT, record_json TEXT);\n')
        for part in self.server.api.iter_batches(a):
            for r in part['batch']['matches']:
                payload=json.dumps(r,ensure_ascii=False).replace("'","''")
                self.wfile.write(("INSERT INTO catalog_records VALUES ('"+part['kind']+"','"+payload+"');\n").encode())
            self.wfile.flush()
        self.wfile.write(b'COMMIT;\n')

    def schema(self,dictionary=False):
        from catalog_queries import ro
        with Catalog(self.server.catalog) as catalog:
            stores={k:v for k,v in catalog.paths.items() if v.suffix=='.sqlite3' and (dictionary or k=='fish_database')}
        out=io.StringIO(newline='')
        if dictionary:
            writer=csv.writer(out);writer.writerow(['store','table','field','type','required','primary_key'])
        for name,path in stores.items():
            conn=ro(path)
            try:
                tables=conn.execute("SELECT name,sql,type FROM sqlite_master WHERE sql IS NOT NULL ORDER BY type,name").fetchall()
                for table,sql,kind in tables:
                    if dictionary and kind=='table':
                        for field in conn.execute('PRAGMA table_info("'+table.replace('"','""')+'")'):
                            writer.writerow([name,table,field[1],field[2],field[3],field[5]])
                    elif not dictionary:out.write('-- '+name+'\n'+sql+';\n\n')
            finally:conn.close()
        self.send(200,('\ufeff'+out.getvalue()).encode(),'text/csv; charset=utf-8' if dictionary else 'application/sql; charset=utf-8',download='data-dictionary.csv' if dictionary else 'schema.sql')
    def log_message(self,*args):pass

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--catalog',type=Path,required=True);p.add_argument('--static',type=Path,default=ROOT/'app/dist');p.add_argument('--port',type=int,default=8771);a=p.parse_args()
    s=ThreadingHTTPServer(('127.0.0.1',a.port),Handler);s.catalog=a.catalog.resolve();s.static=a.static.resolve();s.api=WorkspaceAPI(s.catalog)
    print('http://127.0.0.1:'+str(a.port)+'/?view=1&tab=search',flush=True);s.serve_forever()
